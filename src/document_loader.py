"""
Document Loader Module
Responsible for loading documents (PDF, TXT, DOCX) from disk or in-memory streams,
extracting text, and preserving rich metadata (source filename, page numbers).
"""

import os
from typing import List, Union, BinaryIO
from langchain_core.documents import Document
from pypdf import PdfReader


def load_pdf_from_path(file_path: str) -> List[Document]:
    """Load a PDF from a local file path and return a list of Documents with page metadata."""
    documents = []
    file_name = os.path.basename(file_path)
    with open(file_path, "rb") as f:
        reader = PdfReader(f)
        for idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            text = text.strip()
            if text:
                documents.append(
                    Document(
                        page_content=text,
                        metadata={
                            "source": file_name,
                            "page": idx + 1,
                            "total_pages": len(reader.pages),
                            "file_type": "pdf"
                        }
                    )
                )
    return documents


def load_pdf_from_bytes(file_bytes: BinaryIO, file_name: str) -> List[Document]:
    """Load a PDF from a byte stream (e.g. Streamlit UploadedFile) preserving page numbers."""
    documents = []
    reader = PdfReader(file_bytes)
    total_pages = len(reader.pages)
    for idx, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        text = text.strip()
        if text:
            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": file_name,
                        "page": idx + 1,
                        "total_pages": total_pages,
                        "file_type": "pdf"
                    }
                )
            )
    return documents


def load_txt_from_bytes(file_bytes: BinaryIO, file_name: str) -> List[Document]:
    """Load a plain text file from byte stream."""
    content = file_bytes.read().decode("utf-8", errors="ignore").strip()
    if not content:
        return []
    return [
        Document(
            page_content=content,
            metadata={
                "source": file_name,
                "page": 1,
                "total_pages": 1,
                "file_type": "txt"
            }
        )
    ]


def load_docx_from_bytes(file_bytes: BinaryIO, file_name: str) -> List[Document]:
    """Load a Word document (.docx) from byte stream."""
    try:
        import docx
        doc = docx.Document(file_bytes)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        content = "\n\n".join(paragraphs).strip()
        if not content:
            return []
        return [
            Document(
                page_content=content,
                metadata={
                    "source": file_name,
                    "page": 1,
                    "total_pages": 1,
                    "file_type": "docx"
                }
            )
        ]
    except Exception as e:
        print(f"Error loading docx {file_name}: {e}")
        return []


def load_uploaded_files(uploaded_files) -> List[Document]:
    """
    Ingests a list of Streamlit uploaded files and extracts Documents with metadata.
    Supported types: .pdf, .txt, .docx
    """
    all_documents = []
    for uploaded_file in uploaded_files:
        name = uploaded_file.name
        ext = os.path.splitext(name)[1].lower()
        # Reset file pointer to beginning
        uploaded_file.seek(0)
        
        if ext == ".pdf":
            docs = load_pdf_from_bytes(uploaded_file, name)
            all_documents.extend(docs)
        elif ext == ".txt":
            docs = load_txt_from_bytes(uploaded_file, name)
            all_documents.extend(docs)
        elif ext == ".docx":
            docs = load_docx_from_bytes(uploaded_file, name)
            all_documents.extend(docs)
        else:
            print(f"Unsupported file format: {ext} for file {name}")
            
    return all_documents


def load_directory(directory_path: str) -> List[Document]:
    """Loads all supported documents (.pdf, .txt, .docx) from a directory on disk."""
    documents = []
    if not os.path.exists(directory_path):
        return documents

    for root, _, files in os.walk(directory_path):
        for f in files:
            ext = os.path.splitext(f)[1].lower()
            file_path = os.path.join(root, f)
            if ext == ".pdf":
                documents.extend(load_pdf_from_path(file_path))
            elif ext == ".txt":
                with open(file_path, "rb") as stream:
                    documents.extend(load_txt_from_bytes(stream, f))
            elif ext == ".docx":
                with open(file_path, "rb") as stream:
                    documents.extend(load_docx_from_bytes(stream, f))

    return documents
